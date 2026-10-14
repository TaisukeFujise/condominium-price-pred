#!/usr/bin/env python3
"""template.pptx の 6 枚をそのまま使い、説明の文と表を中身で埋める。

python3 fill_template.py template.pptx out.pptx
"""

import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

FONT = "Noto Sans JP"
EMU = 914400


def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def run(p, text, sz, b=False, color=None, font=True):
    r = sub(p, "a:r")
    rPr = sub(r, "a:rPr", lang="ja-JP", sz=int(sz * 100), b=1 if b else 0)
    if color:
        sub(sub(rPr, "a:solidFill"), "a:srgbClr", val=color)
    if (
        font
    ):  # 本文の枠はレイアウトの書体に任せる。表のセルはテンプレートと同じく直接書く
        for tag in ("a:latin", "a:ea", "a:cs"):
            sub(rPr, tag, typeface=FONT)
    sub(r, "a:t").text = text


BULLETS = {0: ("●", 228600, -228600), 1: ("○", 457200, -228600)}


def fill_body(shape, items, sz0=12, sz1=11, line=None):
    """本文の枠を、(階層, 文, 太字) の並びで埋める。テンプレートと同じ ● ○ の箇条書き。"""
    txBody = shape._element.find(qn("p:txBody"))
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for lvl, text, b in items:
        p = sub(txBody, "a:p")
        if lvl < 0:  # 1文だけの説明は箇条書きにしない
            pPr = sub(p, "a:pPr", marL=0, indent=0, algn="l")
            sub(pPr, "a:buNone")
            run(p, text, sz0, b, font=False)
            continue
        ch, marL, ind = BULLETS[lvl]
        pPr = sub(p, "a:pPr", lvl=lvl, marL=marL, indent=ind, algn="l")
        if (
            line
        ):  # Noto Sans JP は行の高さが大きいので、項目の多いページだけ行間を詰める
            sub(sub(pPr, "a:lnSpc"), "a:spcPct", val=int(line * 1000))
        sub(
            sub(pPr, "a:spcBef"),
            "a:spcPts",
            val=350 if (lvl == 0 and len(txBody) > 3) else 50,
        )
        sub(sub(pPr, "a:spcAft"), "a:spcPts", val=0)
        sub(sub(pPr, "a:buClr"), "a:srgbClr", val="000000" if lvl == 0 else "737373")
        sub(pPr, "a:buSzPct", val=80000 if lvl == 0 else 90000)
        sub(pPr, "a:buFont", typeface="Arial")
        sub(pPr, "a:buChar", char=ch)
        run(p, text, sz0 if lvl == 0 else sz1, b, font=False)


def set_cell(tc, text, sz=10, b=False):
    txBody = tc.find(qn("a:txBody"))
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for line in text.split("\n"):
        p = sub(txBody, "a:p")
        pPr = sub(p, "a:pPr", marL=0, indent=0, algn="l")
        sub(pPr, "a:buNone")
        run(p, line, sz, b)
    tcPr = tc.find(qn("a:tcPr"))
    for k in ("marT", "marB"):
        tcPr.set(k, str(int(0.04 * EMU)))
    tcPr.set("anchor", "ctr")


def fill_table(frame, widths, rows, y, sz=10, head_sz=11, row_h=0.3, bold_first=False):
    """テンプレートの表の列幅を決め、行を足して埋める。見出し行の文言と塗りはテンプレートのまま。"""
    frame.left, frame.top = Inches(0.34), Inches(y)
    tbl = frame._element.find(".//" + qn("a:tbl"))
    for gc, w in zip(
        tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol")), widths, strict=True
    ):
        gc.set("w", str(int(w * EMU)))
    trs = tbl.findall(qn("a:tr"))
    head, blank = trs[0], copy.deepcopy(trs[1])
    for tr in trs[1:]:
        tbl.remove(tr)
    head.set("h", str(int(0.34 * EMU)))
    for tc in head.findall(qn("a:tc")):
        txt = "".join(t.text or "" for t in tc.iter(qn("a:t")))
        set_cell(tc, txt, head_sz, True)
    for cells in rows:
        tr = copy.deepcopy(blank)
        tr.set("h", str(int(row_h * EMU)))
        for i, (tc, text) in enumerate(zip(tr.findall(qn("a:tc")), cells, strict=True)):
            set_cell(tc, text, sz, bold_first and i == 0)
        tbl.append(tr)
    frame.width = Inches(sum(widths))
    frame.height = Inches(0.34 + row_h * len(rows))


def placeholder(slide, idx):
    return next(s for s in slide.placeholders if s.placeholder_format.idx == idx)


prs = Presentation(sys.argv[1])
s1, s2, s3, s4, s5, s6 = prs.slides

# 表紙
for sh in s1.shapes:
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.text = r.text.replace("[グループ名]", "Di molt")

# 役割分担
body = placeholder(s2, 1)
body.top, body.height = Inches(0.98), Inches(4.15)
fill_body(
    body,
    [
        (0, "人の分担", True),
        (
            1,
            "全員：各自のノートブックで EDA を行い、実験ごとに Issue を立てて PR を出した",
            False,
        ),
        (
            1,
            "藤瀬 太翼：実験の土台（構成・時系列 CV・AI への指示ファイル）と開発の手順を整え、損失関数・目的変数の設計などの実験と、PR のレビューを担当した",
            False,
        ),
        (
            1,
            "南條 佑太：isotonic 較正を作り、CV と LB のずれの分析など、PR・Issue での議論とレビューを担った",
            False,
        ),
        (
            1,
            "大久保 翔：最初のベースラインを作り、目的変数の設計の実験を今の土台に移し替えて再実装した",
            False,
        ),
        (
            1,
            "小林 麟太郎：目的変数を総額から㎡単価に変え、間取りを部屋数などに分解する特徴量を試した",
            False,
        ),
        (0, "AIの分担", True),
        (
            1,
            "リポジトリの構成と関数の切り分けの案、前処理・特徴量・学習のコード",
            False,
        ),
        (
            1,
            "Issue にした実験の実行と、結果をまとめた PR の作成、PR の妥当性のレビュー",
            False,
        ),
        (0, "AIの出力の確認", True),
        (
            1,
            "構成とコード：AI と議論し、出力されたコードを VS Code で読みながら直し、ruff の整形と静的チェックの後も、読みにくい箇所や名前は自分で直した",
            False,
        ),
        (
            1,
            "変更の単位：分析の工程ごとにファイルを分け、1 回の PR は 1 つの工程の修正と検証に限って、人が検証結果と実装コードを把握できる範囲に抑えた",
            False,
        ),
        (
            1,
            "検証：AI が出した結果や指摘はそのまま信じず、全員が同じ 3 fold の時系列 CV で測り、提出した LB と連動するかも確かめてから採否を決めた",
            False,
        ),
        (0, "使ったAIツール", True),
        (
            1,
            "Claude Code（藤瀬・南條）、Antigravity CLI（大久保・小林。無料プランで使える範囲が広いため）",
            False,
        ),
    ],
    sz0=11,
    sz1=10,
    line=122,
)

# 試した手法やモデルとその結果
desc = placeholder(s3, 1)
desc.top, desc.height = Inches(0.88), Inches(0.95)
fill_body(
    desc,
    [
        (
            0,
            "評価：RMSLE（小さいほど良い）。過去で学習し、直後の3四半期を当てる 3 fold の時系列 CV の平均を CV スコアとする",
            False,
        ),
        (
            0,
            "表の値：上の行までに採用した変更を積み上げた値",
            False,
        ),
        (
            0,
            "今の最高は LB 0.19709。上位チームは 0.19408 で、まだ 0.003 差がある",
            False,
        ),
    ],
    sz0=9.5,
)
for p in desc._element.iter(qn("a:spcPts")):
    p.set("val", "50")
frame = next(sh for sh in s3.shapes if sh.has_table)
fill_table(
    frame,
    [4.0, 1.35, 1.1, 2.87],
    [
        ["LightGBM のベースライン", "0.20674", "0.20542", "基準"],
        [
            "cat_smooth（カテゴリの扱い）と learning_rate（学習率）の調整",
            "0.20360",
            "0.20222",
            "採用。CV と LB が同じ幅で改善",
        ],
        [
            "損失関数を二乗誤差から Huber 損失に変更",
            "0.20004",
            "0.19914",
            "採用。3 fold すべてで改善",
        ],
        [
            "isotonic 較正で予測を後から補正",
            "0.19727",
            "0.20059",
            "不採用。CV は良いが LB で悪化した",
        ],
        [
            "目的変数を㎡単価にする（予測に面積を掛けて戻す）",
            "0.19845",
            "—",
            "採用",
        ],
        [
            "目的変数の設計を見直す",
            "0.19664",
            "0.19709",
            "採用予定（レビュー中）。LB が最高",
        ],
        [
            "効かなかったもの：場所の特徴量・容積率の補完など（Huber 時点）、外れ値の重み・間取りの分解（最新の土台の上）",
            "差 0.0004 以内",
            "—",
            "不採用",
        ],
    ],
    y=1.85,
    sz=8.5,
    head_sz=9.5,
    row_h=0.36,
)

# AIの提案：採用したもの・却下したもの
desc = placeholder(s4, 1)
fill_body(
    desc,
    [
        (
            -1,
            "構成と関数の分け方は AI の案を議論して直してから採用し、実験の案・集計・レビューは CV とデータで確かめて採否を決めた",
            False,
        ),
    ],
    sz0=11,
)
frame = next(sh for sh in s4.shapes if sh.has_table)
fill_table(
    frame,
    [3.4, 1.65, 4.27],
    [
        [
            "リポジトリ設計・コード：実験ごとに記録用のフォルダを作り、工程番号つきのファイルに分ける。学習と予測を 1 つの関数で行う",
            "修正して採用",
            "議論して、実験の記録は PR に一本化し、ファイルは工程ごと、関数は CV 用と本番用に分けた。分ける前後で CV が完全に一致することを確かめた",
        ],
        [
            "実験の案：件数の少ないカテゴリでも木が分岐できるよう、パラメータを下げる",
            "却下",
            "CV が変わらず、逆向き（件数の少ないカテゴリを全体の平均に寄せる）のほうが効いた",
        ],
        [
            "集計：価格のばらつきのうち、容積率で説明できる割合は 0.6%",
            "却下（数値を訂正）",
            "条件を変えて数え直し、値が1種類しかない地区まで混ぜていた誤りを見つけた（正しくは 0.9%）",
        ],
        [
            "PR のレビュー：較正の移植で見つけた不具合 2 件",
            "採用",
            "直したあとも CV が変わらないことを確かめた",
        ],
        [
            "PR のレビュー：分位点の取り方と、欠けた値の埋め方がおかしい",
            "却下",
            "データで確かめ、前者はわざとそうしていて、後者は今のデータでは起きないと分かった",
        ],
        [
            "実装：理由の書かれていない決め打ちの値（期間の区切り・グループの数・くり返しの回数）",
            "修正して採用",
            "値を変えて結果への影響を確かめ、定数にして理由を書くよう直した。影響の大きい値は今後の改善の対象にした",
        ],
    ],
    y=1.55,
    sz=9.5,
    head_sz=10,
    row_h=0.5,
)

# 共同開発方法
body = placeholder(s5, 1)
body.top, body.height = Inches(0.98), Inches(4.15)
fill_body(
    body,
    [
        (0, "リポジトリの構成", True),
        (
            1,
            "main.py は工程を順に呼ぶだけにし、前処理・特徴量・目的変数・モデルなど、工程ごとにファイルを 1 つにした",
            False,
        ),
        (
            1,
            "特徴量は MAKERS、モデルは MODELS に 1 行足すだけで試せる。別々の実験を最新の土台に積み直しても衝突しなかった",
            False,
        ),
        (0, "実験の進め方（GitHub）", True),
        (
            1,
            "1 つの実験を Issue → ブランチ → PR で進め、PR に fold ごとの CV・比べた main のコミット・Public LB を書く",
            False,
        ),
        (
            1,
            "効かなかった実験も「不採用」ラベルで閉じて残し、同じ実験を二度試さない",
            False,
        ),
        (
            1,
            "レビューは藤瀬が中心に行い、AI のレビューの指摘もデータで確かめてから直してマージする",
            False,
        ),
        (0, "AIへの指示ファイル（CLAUDE.md・AGENTS.md）", True),
        (
            1,
            "テスト期間（2025年第3四半期以降）の取引データを取得・参照しない規約と、検証方法を合意なしに変えないこと",
            False,
        ),
        (
            1,
            "価格を使う処理は fold ごとに作ること、型ヒント・docstring・名前の決まり",
            False,
        ),
        (0, "グループ内の情報共有", True),
        (
            1,
            "Discord で PR を報告して全員が確認し、submission チャンネルで日ごとの提出担当を割り当てる",
            False,
        ),
        (1, "10/11 に全員で MTG を開き、土台の構成と担当を決めた", False),
    ],
    sz0=11.5,
    sz1=10.5,
)

# 良かった点・悪かった点
fill_body(
    placeholder(s6, 1),
    [
        (0, "良かった点", True),
        (
            1,
            "チームの土台の上では全員が同じ時系列 CV で比べたので、施策の良し悪しを同じ基準で判断できた",
            False,
        ),
        (
            1,
            "CV と LB のずれを提出で切り分け、isotonic 較正が test で効いていないと分かって外せた",
            False,
        ),
        (
            1,
            "AI のレビューの指摘も、設計の元にした EDA の値も、データで確かめてから採否を決めた",
            False,
        ),
        (1, "不採用の実験も PR に残し、同じ実験を二度試さずに済んだ", False),
        (0, "悪かった点", True),
        (1, "土台の整備に時間がかかり、実験の着手が遅れた", False),
        (
            1,
            "isotonic 較正を CV だけで「採用予定」にし、LB で悪化すると分かるまで時間がかかった",
            False,
        ),
        (
            1,
            "EDA で決めた値の一部に根拠がなかった（結果への影響はほぼなし）",
            False,
        ),
        (0, "最終報告までに変えること", True),
        (
            1,
            "提出の担当者が自分で main を動かして提出し、コードの理解とコミットが1人に偏らないようにする",
            False,
        ),
        (
            1,
            "効いた特徴量を積み上げ、モデルを混ぜる",
            False,
        ),
    ],
    sz0=12,
    sz1=11,
)

prs.save(sys.argv[2])
print("saved", sys.argv[2])
